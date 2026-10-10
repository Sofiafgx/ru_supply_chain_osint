"""Offline customs imports with exact, auditable employer-name fallback."""
import argparse
import csv
import hashlib
import json
import sqlite3
import unicodedata
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from trade_interval_cli import FIELDS, normalize, text


def valid_inn(value):
    value = text(value)
    return bool(value and re.fullmatch(r'\d{10}', value) and
                sum(int(x)*w for x, w in zip(value[:9], (2,4,10,3,5,9,4,6,8))) % 11 % 10 == int(value[-1]))


def normalized_name(value):
    value = unicodedata.normalize('NFKC', text(value) or '').casefold().replace('ё', 'е')
    words = re.findall(r'[^\W_]+', value, re.UNICODE)
    # Remove leading legal form only; retain branch, city and substantive name tokens.
    while words and words[0] in ('ао', 'пао', 'оао', 'зао', 'ооо'):
        words.pop(0)
    return ' '.join(words)


def employer_index(path):
    index = defaultdict(list)
    with path.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        if not {'name', 'inn'}.issubset(reader.fieldnames or []):
            raise ValueError('Employer dictionary requires name and inn columns')
        for row in reader:
            key = normalized_name(row['name'])
            if key:
                index[key].append({k: row.get(k) for k in ('employer_id', 'name', 'inn', 'match_method', 'profile_url')})
    return index


def resolve_importer(row, index):
    resolution = {'original_importer_tax_id': row.get('importer_tax_id'),
                  'original_importer_name': row.get('importer_name')}
    if text(row.get('importer_tax_id')) is not None:
        resolution['method'] = 'provided_inn' if valid_inn(row['importer_tax_id']) else 'invalid_provided_inn'
        return row, resolution
    key = normalized_name(row.get('importer_name'))
    candidates = index.get(key, []) if key else []
    resolution.update(normalized_name=key, dictionary_candidates=candidates)
    ids = {text(r['inn']) for r in candidates}
    if len(ids) == 1 and valid_inn(next(iter(ids))):
        row = dict(row, importer_tax_id=next(iter(ids)))
        resolution.update(method='unique_exact_normalized_name', resolved_inn=row['importer_tax_id'],
                          identity_basis='employer_dictionary_inference; not a customs-provided INN')
    else:
        resolution['method'] = 'ambiguous_or_invalid_dictionary_match' if candidates else 'no_dictionary_match'
    return row, resolution


def load_local(path, config):
    suffix = path.suffix.lower()
    if suffix == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            reader = csv.DictReader(handle, delimiter=config.get('csv_delimiter', ','))
            return list(reader)
    if suffix in ('.json', '.jsonl', '.ndjson'):
        content = path.read_text(encoding='utf-8-sig')
        return [json.loads(line) for line in content.splitlines() if line.strip()] if suffix != '.json' else json.loads(content)
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[config['sheet']] if config.get('sheet') else workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator)
            if None in headers or len(set(headers)) != len(headers):
                raise ValueError('XLSX header must contain unique non-empty column names')
            return [dict(zip(headers, values)) for values in iterator if any(v is not None for v in values)]
        finally:
            workbook.close()
    if suffix == '.parquet':
        import pyarrow.parquet as parquet
        return parquet.read_table(path).to_pylist()
    if suffix in ('.sqlite', '.sqlite3', '.db'):
        table = config.get('sqlite_table')
        if not table:
            raise ValueError('sqlite_table must explicitly identify the local customs table')
        with sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True) as connection:
            connection.row_factory = sqlite3.Row
            if not connection.execute('SELECT 1 FROM sqlite_master WHERE type=? AND name=?', ('table', table)).fetchone():
                raise ValueError('Customs table not found')
            return [dict(row) for row in connection.execute('SELECT * FROM "' + table.replace('"', '""') + '"')]
    raise ValueError('Supported: CSV, JSON array, JSONL, XLSX, Parquet, SQLite. Other dump formats require conversion.')


def run(config_path, input_override=None, output_override=None):
    base = config_path.resolve().parent
    config = json.loads(config_path.read_text(encoding='utf-8'))
    def path(value):
        p = Path(value)
        return p if p.is_absolute() else base / p
    supplied = input_override or config.get('input_file')
    if not supplied:
        raise ValueError('Provide --input or set input_file to a real local customs export')
    input_path = path(supplied)
    output = path(output_override or config['output_dir'])
    dictionary = path(config['employers_file'])
    plan_path = path(config['plan_file'])
    plan = json.loads(plan_path.read_text())
    rows = load_local(input_path, config)
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError('Expected a JSON array or tabular row objects')
    index = employer_index(dictionary)
    processed, resolutions = [], []
    allowed = set(FIELDS) | {'importer_name', 'import_country', 'source', 'source_url', 'date_type',
                            'quantity', 'quantity_unit', 'value_type', 'value_currency', 'shipment_id'}
    mapping = config.get('column_map', {})
    if set(mapping) - allowed or set(config.get('defaults', {})) - allowed:
        raise ValueError('Unknown canonical field in column_map or defaults')
    if rows:
        missing_columns = set(mapping.values()) - set().union(*(set(row) for row in rows))
        if missing_columns:
            raise ValueError('Mapped columns not found: ' + ', '.join(sorted(missing_columns)))
    for raw in rows:
        row = dict(config.get('defaults', {}))
        # Explicit mappings are applied to each whole row, not guessed from labels.
        for canonical in allowed:
            column = mapping.get(canonical, canonical)
            if column in raw and text(raw[column]) is not None:
                value = raw[column]
                row[canonical] = value.isoformat()[:10] if isinstance(value, (datetime, date)) else value
        row, resolution = resolve_importer(row, index)
        processed.append(row)
        resolutions.append(resolution)
    records, audit, summaries = normalize(processed, plan)
    for item in audit:
        item['identity_resolution'] = resolutions[item['input_row'] - 1]
        item['input_raw'] = rows[item['input_row'] - 1]
    # Check all outputs before touching input or reference files.
    protected = {input_path.resolve(), dictionary.resolve(), plan_path.resolve(), config_path.resolve()}
    if any((output / name).resolve() in protected for name in
           ('shipments.json', 'shipments.csv', 'row_audit.json', 'rhythm_summary.json', 'run_manifest.json')):
        raise ValueError('Output would overwrite input or reference data')
    output.mkdir(parents=True, exist_ok=True)
    def digest(p):
        with p.open('rb') as handle:
            h = hashlib.sha256()
            for chunk in iter(lambda: handle.read(1024*1024), b''):
                h.update(chunk)
            return h.hexdigest()
    manifest = {'mode': 'offline', 'input_file': str(input_path.resolve()), 'input_sha256': digest(input_path),
                'employers_file': str(dictionary.resolve()), 'employers_sha256': digest(dictionary),
                'plan_sha256': digest(plan_path), 'config_sha256': digest(config_path),
                'input_rows': len(rows), 'normalized_rows': len(records),
                'name_fallback_rows': sum(r['method'] == 'unique_exact_normalized_name' for r in resolutions),
                'interval_grouping': 'source + importer INN + HS6; all exporters',
                'coverage': 'provided_export_only'}
    for name, data in [('shipments.json', records), ('row_audit.json', audit),
                       ('rhythm_summary.json', summaries), ('run_manifest.json', manifest)]:
        (output / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + '\n')
    with (output / 'shipments.csv').open('w', encoding='utf-8', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(records)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', type=Path, default=Path(__file__).resolve().parents[1] / 'config/local_trade_pipeline.json')
    parser.add_argument('--input', help='Real local export; relative paths resolve against config directory')
    parser.add_argument('--output-dir')
    args = parser.parse_args()
    try:
        print(json.dumps(run(args.config, args.input, args.output_dir), ensure_ascii=False))
    except (ValueError, OSError, ImportError) as exc:
        parser.exit(1, str(exc) + '\n')


if __name__ == '__main__':
    main()
